"""A chargeback rebuttal, from the dispute facts and what the board remembers.

The facts here are Jose Zambrano's and Skip Scott's real disputes.
"""

from __future__ import annotations

from datetime import date

import pytest

from wilbyte import rebuttal

JOSE = """MID: 510200014664

DBA Name: AGENT LEAD LAB

Dispute Date: 9/8/2026

Dispute Type: Debited

Dispute Dollar Amount: $1,552.50

Acquirer's Reference Number: 24556406167808942703416

Card Number: ending in 2610

Transaction Date: 6/15/2026

Customer Name: Jose Zambrano

Customer Email: josezagent@gmail.com
"""


# ------------------------------------------------------------- reading it


def test_the_whole_block_is_read():
    one = rebuttal.read_facts(JOSE)

    assert one.mid == "510200014664"
    assert one.dba == "AGENT LEAD LAB"
    assert one.dispute_type == "Debited"
    assert one.arn == "24556406167808942703416"
    assert one.card == "ending in 2610"
    assert one.customer_name == "Jose Zambrano"
    assert one.customer_email == "josezagent@gmail.com"


def test_the_amount_is_normalised():
    assert rebuttal.read_facts(JOSE).amount == "$1,552.50"


@pytest.mark.parametrize(
    "said, expected",
    [("6/15/2026", date(2026, 6, 15)), ("06/15/26", date(2026, 6, 15)),
     ("2026-06-15", date(2026, 6, 15)), ("not a date", None), ("", None)],
)
def test_a_date_is_read_however_it_was_typed(said, expected):
    assert rebuttal.as_date(said) == expected


def test_a_label_he_does_not_know_is_ignored_rather_than_guessed_at():
    """A field RYTE invented is worse than one left blank."""
    one = rebuttal.read_facts("Something Else: 12345\nCustomer Name: Jose Zambrano")

    assert one.customer_name == "Jose Zambrano"
    assert one.mid == ""


def test_the_first_label_that_fits_wins():
    """"Customer Email" and "Email" both match the email field; the block is
    read top down and the more specific one is written first."""
    one = rebuttal.read_facts(
        "Customer Email: jose@example.com\nEmail: someone-else@example.com"
    )

    assert one.customer_email == "jose@example.com"


def test_a_block_missing_what_matters_says_which():
    one = rebuttal.read_facts("MID: 510200014664\nDBA Name: AGENT LEAD LAB")

    assert one.missing() == [
        "Customer Name", "Dispute Dollar Amount", "Transaction Date",
    ]


def test_a_full_block_is_missing_nothing():
    assert rebuttal.read_facts(JOSE).missing() == []


def test_the_name_can_come_off_the_command_line_instead():
    assert rebuttal.named_in("rebuttal Jose Zambrano") == "Jose Zambrano"
    assert rebuttal.named_in("chargeback Skip Scott\nMID: 1") == "Skip Scott"


# ----------------------------------------------------------- how long they waited


def test_the_gap_between_paying_and_disputing():
    """85 days. Somebody who received nothing does not wait three months."""
    assert rebuttal.read_facts(JOSE).days_waited() == 85


def test_the_gap_is_said_out_loud():
    said = rebuttal.waited_line(rebuttal.read_facts(JOSE))

    assert "85 days" in said
    assert "6/15/2026" in said and "9/8/2026" in said


def test_a_dispute_filed_the_same_week_makes_no_such_argument():
    one = rebuttal.read_facts(
        "Customer Name: A\nDispute Dollar Amount: $1\n"
        "Transaction Date: 6/15/2026\nDispute Date: 6/18/2026"
    )

    assert one.days_waited() == 3
    assert rebuttal.waited_line(one) == ""


def test_no_dispute_date_means_no_gap_and_no_line():
    one = rebuttal.read_facts(
        "Customer Name: A\nDispute Dollar Amount: $1\nTransaction Date: 6/15/2026"
    )

    assert one.days_waited() is None
    assert rebuttal.waited_line(one) == ""


# ------------------------------------------------------------------ the header


def test_a_blank_field_is_left_out_of_the_fact_table():
    rows = rebuttal.header(rebuttal.read_facts(
        "Customer Name: Jose Zambrano\nDispute Dollar Amount: $1,552.50\n"
        "Transaction Date: 6/15/2026"
    ))
    labels = [label for label, _value in rows]

    assert "Acquirer Reference Number" not in labels
    assert "Amount" in labels


def test_the_email_rides_with_the_name():
    rows = dict(rebuttal.header(rebuttal.read_facts(JOSE)))

    assert "Jose Zambrano" in rows["Cardholder"]
    assert "josezagent@gmail.com" in rows["Cardholder"]


def test_the_dates_are_written_out_rather_than_slashed():
    """A slashed date in a legal document reads as a form somebody filled in."""
    rows = dict(rebuttal.header(rebuttal.read_facts(JOSE)))

    assert rows["Transaction Date"] == "June 15, 2026"
    assert rows["Dispute Date"] == "September 8, 2026"


# ------------------------------------------------------------------- exhibits


def test_a_file_knows_what_sort_of_file_it_is():
    assert rebuttal.Exhibit("texts.PNG").is_image() is True
    assert rebuttal.Exhibit("contract.pdf").is_pdf() is True
    assert rebuttal.Exhibit("notes.txt").is_image() is False


def test_every_exhibit_kind_has_a_proof_to_sit_under():
    for kind in rebuttal.EXHIBITS:
        assert kind == "other" or kind in rebuttal.UNDER


# --------------------------------------------------------------- what is missing


def test_a_proof_with_nothing_behind_it_is_named_not_written_around():
    """A gap named is a gap somebody fills. A gap written around reaches the
    acquirer."""
    holes = rebuttal.what_is_missing(rebuttal.Gathered(), [])

    assert any("PandaDoc" in one for one in holes)
    assert any("invoice" in one.casefold() for one in holes)
    assert any("texts" in one.casefold() for one in holes)


def test_an_attached_contract_closes_that_hole():
    holes = rebuttal.what_is_missing(
        rebuttal.Gathered(), [rebuttal.Exhibit("c.pdf", kind="contract")]
    )

    assert not any("PandaDoc" in one for one in holes)


def test_the_holes_the_gathering_found_are_carried_through():
    found = rebuttal.Gathered(holes=["No sheet link on their card"])

    assert "No sheet link on their card" in rebuttal.what_is_missing(found, [])


# -------------------------------------------------------------- what it says


def test_the_writing_prompt_carries_only_what_was_gathered():
    """Nothing is invented, so nothing that was not found is described."""
    found = rebuttal.Gathered(delivery="Faith: your launch date is Monday")

    asked = rebuttal.writing_prompt(rebuttal.read_facts(JOSE), found, [])

    assert "onboarding and delivery" in asked
    assert "### the delivered lead sheet" not in asked
    assert "cite only what is above" in asked.casefold()


def test_the_demand_names_the_amount():
    assert "$1,552.50" in rebuttal.demand(rebuttal.read_facts(JOSE))


def test_the_demand_still_reads_without_one():
    assert "the disputed amount" in rebuttal.demand(rebuttal.Dispute())


# --------------------------------------------------- getting to the command


def test_a_dispute_block_is_not_a_request_for_an_email():
    """The block ends "Customer Email: ...", and "email" is a copy format —
    so the first rebuttal ever asked for came back as a marketing email about
    a chargeback."""
    from wilbyte.bot import mentions

    asked = mentions.parse("<@1> " + JOSE.replace("MID:", "rebuttal\nMID:", 1))

    assert asked.action == "rebuttal"
    assert asked.format_key is None


@pytest.mark.parametrize("word", ["rebuttal", "chargeback", "dispute"])
def test_every_way_of_asking_reaches_it(word):
    from wilbyte.bot import mentions

    asked = mentions.parse(f"<@1> {word} Jose Zambrano\nCustomer Email: a@b.com")

    assert asked.action == "rebuttal"


def test_the_whole_block_travels_with_it():
    """Not the remainder — every field is needed and they are on their own
    lines, so nothing may be stripped off the front."""
    from wilbyte.bot import mentions

    asked = mentions.parse("<@1> rebuttal\n" + JOSE)

    assert "24556406167808942703416" in (asked.brief or "")
    assert "josezagent@gmail.com" in (asked.brief or "")


@pytest.mark.parametrize(
    "said, expected",
    [("email about the new aged lead prices", "email"),
     ("sms for the OTP launch", "sms"),
     ("ad for agents stuck at 20 leads a week", "ad")],
)
def test_asking_for_copy_still_writes_copy(said, expected):
    """The fix must not cost the thing it was checked before."""
    from wilbyte.bot import mentions

    asked = mentions.parse(f"<@1> {said}")

    assert asked.action == "write"
    assert asked.format_key == expected


# ------------------------------------------- when it goes wrong, saying so


def test_the_evidence_gathering_reads_their_card_against_the_day_they_paid():
    """Their launch date is months back by the time a dispute lands. Read
    against today, "Monday" is next Monday."""
    import inspect

    from wilbyte.bot import jobs

    source = inspect.getsource(jobs.rebuttal_evidence)

    assert "read_agent(" in source
    assert "today=" in source.split("read_agent(")[1][:300]


def test_the_failure_message_can_actually_be_sent(monkeypatch):
    """The first real run failed on a missing argument, and the handler that
    was supposed to say so raised NameError itself — so the error nobody could
    see was replaced by a generic one. An error path that cannot run is worse
    than no error path."""
    import asyncio

    from wilbyte.bot import client as bot_client
    from wilbyte.bot import jobs

    said = []

    class Responder:
        requester_id = 1

        async def send(self, content=None, **kwargs):
            said.append(str(content or kwargs.get("embed") or ""))

    class Message:
        attachments: list = []

    def boom(*args, **kwargs):
        raise TypeError("read_agent() missing 1 required keyword-only argument")

    monkeypatch.setattr(jobs, "rebuttal_evidence", boom)
    monkeypatch.setattr(
        bot_client.embeds, "error", lambda text, **kw: f"ERROR: {text}"
    )

    class Config:
        class secrets:
            anthropic_api_key = "x"

    asyncio.run(bot_client._rebuttal(Responder(), Config(), Message(), JOSE))

    assert any("read_agent" in one for one in said), said


# ----------------------------------------------- running it again from a reply


@pytest.mark.parametrize(
    "content, expected",
    [("<@1>", False), ("<@1>  ", False), ("<@1> <@&2>", False),
     ("<@1> rebuttal", True), ("<@1> hi", True)],
)
def test_a_bare_mention_is_recognised_as_saying_nothing(content, expected):
    from wilbyte.bot import mentions

    assert mentions.said_anything(content) is expected


def test_replying_to_a_command_and_saying_nothing_runs_it_again(monkeypatch):
    """Franklin pasted the dispute notice with six screenshots on it, then
    replied to that message with "@Ryte" three times and got the help card.
    Nobody is going to paste a dispute notice twice."""
    import asyncio

    from wilbyte.bot import client as bot_client
    from wilbyte.bot import mentions

    ran = []

    class Older:
        content = "<@1> rebuttal\n" + JOSE
        attachments = ["a screenshot"]

    class Bare:
        content = "<@1>"
        attachments: list = []
        author = type("A", (), {"id": 2, "bot": False})()

    async def replied_to(_message):
        return Older()

    async def rebuttal_ran(responder, config, message, said):
        ran.append((message, said))

    monkeypatch.setattr(bot_client, "_replied_to", replied_to)
    monkeypatch.setattr(bot_client, "_rebuttal", rebuttal_ran)

    asked = mentions.parse(Bare.content)
    assert asked.action == "help"  # on its own, it is nothing

    again = mentions.parse(Older.content)
    assert again.action == "rebuttal"  # ...and the one it answers is the command


def test_the_attachments_come_from_the_message_that_had_them():
    """The screenshots were on the original. A reply carries none of them."""
    import inspect

    from wilbyte.bot import client as bot_client

    source = inspect.getsource(bot_client.handle_mention)

    assert "_replied_to" in source
    assert "message = replied" in source


# ---------------------------------------------------------------- exhibits at the back


def test_attachments_are_grouped_and_lettered_in_the_order_they_are_argued():
    """The messages before the purchase come before the contract, which comes
    before the invoice, which comes before what was delivered."""
    ex = [
        rebuttal.Exhibit("sheet.png", kind="sheet"),
        rebuttal.Exhibit("chat1.png", kind="texts"),
        rebuttal.Exhibit("deal.pdf", kind="contract"),
        rebuttal.Exhibit("chat2.png", kind="texts"),
    ]

    lettered = rebuttal.letter_them(ex)

    assert [(one.name, one.label()) for one in lettered] == [
        ("chat1.png", "Exhibit A — screenshot 1 of 2"),
        ("chat2.png", "Exhibit A — screenshot 2 of 2"),
        ("deal.pdf", "Exhibit B"),
        ("sheet.png", "Exhibit C"),
    ]


def test_one_of_a_kind_is_not_called_screenshot_one_of_one():
    (one,) = rebuttal.letter_them([rebuttal.Exhibit("i.pdf", kind="invoice")])

    assert one.label() == "Exhibit A"


def test_the_groups_carry_a_name_a_reader_understands():
    ex = rebuttal.letter_them([
        rebuttal.Exhibit("c.png", kind="texts"),
        rebuttal.Exhibit("d.pdf", kind="invoice"),
    ])

    assert [(letter, what) for letter, what, _group in rebuttal.exhibit_groups(ex)] == [
        ("A", "WhatsApp / Text Conversation with the Cardholder"),
        ("B", "Invoice and Payment"),
    ]


def test_nothing_attached_is_no_exhibits_rather_than_an_empty_heading():
    assert rebuttal.exhibit_groups([]) == []


def test_the_prompt_gives_claude_the_translations_to_argue_from():
    """The exhibits reach the writing as what they say, not as filenames —
    otherwise it can cite Exhibit A without knowing what is in it."""
    ex = rebuttal.letter_them([
        rebuttal.Exhibit(
            "c.png", kind="texts",
            transcript="8:37 PM Cardholder: I can pay by card, right?",
        )
    ])

    asked = rebuttal.writing_prompt(rebuttal.read_facts(JOSE), rebuttal.Gathered(), ex)

    assert "EXHIBIT A" in asked
    assert "I can pay by card" in asked
    assert "Cite exhibits by letter, and only ones that exist." in asked


# ------------------------------------------------------- the shape of the page


def test_a_conversation_is_put_back_into_the_order_it_happened():
    """WhatsApp hands its screenshots over newest first, so the first one Jose
    attached was the last thing said. The rebuttal read backwards."""
    ex = [
        rebuttal.Exhibit("c.png", kind="texts", transcript="— Wed, Jun 17 —\n10:40 AM"),
        rebuttal.Exhibit("a.png", kind="texts", transcript="— Fri, Jun 12 —\n8:04 PM"),
        rebuttal.Exhibit("b.png", kind="texts", transcript="— Sun, Jun 14 —\n3:12 PM"),
    ]

    assert [one.name for one in rebuttal.letter_them(ex)] == ["a.png", "b.png", "c.png"]


def test_one_it_cannot_date_keeps_its_place_rather_than_being_guessed_at():
    ex = [
        rebuttal.Exhibit("a.png", kind="texts", transcript="— Fri, Jun 12 —"),
        rebuttal.Exhibit("undated.png", kind="texts", transcript="no date in this one"),
    ]

    assert [one.name for one in rebuttal.letter_them(ex)] == ["a.png", "undated.png"]


def test_only_conversations_are_reordered():
    """Two pages of a contract are in the order they were attached, and that
    is the order they belong in."""
    ex = [
        rebuttal.Exhibit("p2.png", kind="contract", transcript="— Jun 17 —"),
        rebuttal.Exhibit("p1.png", kind="contract", transcript="— Jun 12 —"),
    ]

    assert [one.name for one in rebuttal.letter_them(ex)] == ["p2.png", "p1.png"]


@pytest.mark.parametrize(
    "said", ["— Wed, Jun 17 —", "Jun 17, 2026", "2026-06-17", "June 17"]
)
def test_a_date_is_found_however_the_transcript_writes_it(said):
    assert rebuttal.happened_on(said)[1:] == (6, 17)


def test_the_writing_is_asked_for_markers_rather_than_a_shape_to_guess_at():
    """The first version guessed a heading from "a numbered line with no full
    stop", and the writing came back numbered with full stops — so every
    argument rendered as a bullet and the document had no structure at all."""
    asked = rebuttal.writing_prompt(rebuttal.read_facts(JOSE), rebuttal.Gathered(), [])

    assert "ARGUMENT:" in asked
    assert "Do not number them" in asked
    assert "no full stop at the end" in asked


# ------------------------------------------------------------- key messages


def test_the_messages_are_pulled_out_to_be_set_as_a_table():
    from wilbyte.bot import jobs

    written = jobs._split_written(
        "SUMMARY:\nHe bought leads.\n\n"
        "ARGUMENT: He chose the tier in writing (Exhibit A)\nOn June 14 he said so.\n\n"
        "KEY MESSAGES:\n"
        'Jun 12, 8:37 PM | Cardholder | "I can pay by card, right?"\n'
        'Jun 14, 3:12 PM | Cardholder | "To start yes, I want to try them"\n\n'
        "CONCLUSION:\nReverse it."
    )

    assert "I can pay by card" in written["messages"]
    assert "KEY MESSAGES" not in written["body"]
    assert "CONCLUSION" in written["body"]
    assert "ARGUMENT" in written["body"]


def test_no_message_block_leaves_the_writing_alone():
    from wilbyte.bot import jobs

    said = "SUMMARY:\nHe bought leads.\n\nCONCLUSION:\nReverse it."

    assert jobs._split_written(said) == {"body": said}


def test_the_messages_sit_under_the_summary_not_after_the_conclusion():
    """They are what the arguments are about, and an acquirer who reads only
    the first page should be reading them."""
    from wilbyte import rebuttaldoc

    opening, arguments = rebuttaldoc._split_at_first_argument(
        "SUMMARY:\nHe bought leads.\n\nARGUMENT: He chose it (Exhibit A)\nHe did."
    )

    assert opening.startswith("SUMMARY:")
    assert arguments.startswith("ARGUMENT:")


def test_a_body_with_no_arguments_is_not_cut_in_half():
    from wilbyte import rebuttaldoc

    opening, arguments = rebuttaldoc._split_at_first_argument("SUMMARY:\nJust this.")

    assert opening == "SUMMARY:\nJust this."
    assert arguments == ""


def test_the_writing_is_asked_for_the_messages_and_told_when_to_leave_them_out():
    asked = rebuttal.writing_prompt(rebuttal.read_facts(JOSE), rebuttal.Gathered(), [])

    assert "KEY MESSAGES:" in asked
    assert "when | who | what they said" in asked
    assert "Leave this out entirely if there are no messages" in asked


# ------------------------------- a client who is not a new agent is still a client

# Juliana Hernandez's real dispute. Her card was titled "AGED LEAD - Juliana
# Hernandez" — an order from somebody already a client rather than a fresh
# setup — so a search for New Agent cards reported that the board had never
# heard of her, and the rebuttal went out asking Franklin to go and find the
# order, the sheet and the invoice that RYTE could already see.

JULIANA = """Disputed Payment | Elevateqs

ARN: 72307626241809574244780
Customer Name: Juliana Hernandez
Customer Email: hjuliana650@gmail.com
Card Number (Last 4): 7543
Transaction Date: 8/28/2026
Dispute Amount: $ 129.37"""

HER_CARD = {"id": "c1", "name": "AGED LEAD - Juliana Hernandez"}

HER_DESC = """Name: Juliana Hernandez
Email: hjuliana650@gmail.com
Phone Number: +14302302708
Lead Type: Aged Final Expense - 30-90 days
States Leads In: Texas
Number of Requested Leads: 25"""

HER_COMMENTS = [
    "delivered",
    "https://docs.google.com/spreadsheets/d/1bX2y6jK_JohC3bjuf/edit?usp=sharing",
    "takeover financial",
]


def test_an_aged_lead_card_is_a_card_about_that_client():
    from wilbyte import agents

    assert agents.named_that("Juliana Hernandez", [HER_CARD]) == [HER_CARD]
    assert agents.agent_name(HER_CARD["name"]) == "Juliana Hernandez"


def test_an_aged_lead_order_is_still_not_a_new_agent_to_be_set_up():
    """Widening the *lookup* must not widen what gets filed as a new setup."""
    from wilbyte import agents

    assert agents.is_agent_card(HER_CARD["name"]) is False
    assert agents.is_client_card(HER_CARD["name"]) is True


def test_a_setup_card_is_preferred_when_a_client_has_both():
    from wilbyte import agents

    found = agents.named_that("Juliana Hernandez", [
        HER_CARD, {"id": "c2", "name": "New Agent - Juliana Hernandez"},
    ])

    assert [one["id"] for one in found] == ["c2", "c1"]


def test_a_daily_card_is_still_not_anybody(monkeypatch):
    from wilbyte import agents

    assert agents.named_that("Lead Order", [{"name": "Lead Order 08/28"}]) == []


def _gathering(monkeypatch, *, cards, receipt=("Reference # FJZ3FV3XAAJF-PXX9", "")):
    """rebuttal_evidence with Trello, Sheets and Gmail all stubbed."""
    from types import SimpleNamespace

    from wilbyte import rebuttal
    from wilbyte.bot import jobs

    class Board:
        def board_cards(self, board_id, archived=False):
            return list(cards)

        def card_detail(self, card_id):
            return {"desc": HER_DESC}

        def card_comments(self, card_id):
            return list(HER_COMMENTS)

        def card_notes(self, card_id):
            return [
                {"text": one, "when": "2026-08-28T18:10:00.000Z", "id": str(i)}
                for i, one in enumerate(HER_COMMENTS)
            ]

        def close(self):
            pass

    monkeypatch.setattr(jobs, "open_trello", lambda config: Board())
    monkeypatch.setattr(jobs, "_payment_receipt", lambda config, dispute: receipt)
    monkeypatch.setattr(
        jobs, "_read_lead_sheet", lambda config, link, gs: ("25 rows of leads", "")
    )

    config = SimpleNamespace(secrets=SimpleNamespace(trello_board_id="b1"))
    return jobs.rebuttal_evidence(config, rebuttal.read_facts(JULIANA))


def test_her_order_and_her_sheet_are_found_off_the_aged_lead_card(monkeypatch):
    found = _gathering(monkeypatch, cards=[HER_CARD])

    assert "Aged Final Expense - 30-90 days" in found.invoice
    assert "Number of Requested Leads: 25" in found.invoice
    assert found.sheet == "25 rows of leads"
    assert "delivered" in found.delivery
    assert not any("anywhere on the board" in one for one in found.holes)


def test_no_card_at_all_still_fetches_the_paid_invoice(monkeypatch):
    """The receipt is in Gmail and has nothing to do with the board.

    Returning early on a missing card meant one silent skip took the invoice
    with it — and the rebuttal asked Franklin to go and find a receipt RYTE
    was already able to read.
    """
    found = _gathering(monkeypatch, cards=[])

    assert "FJZ3FV3XAAJF-PXX9" in found.invoice, "the receipt was dropped again"
    assert any("anywhere on the board" in one for one in found.holes)


def test_no_card_does_not_also_complain_about_a_missing_sheet_link(monkeypatch):
    """One hole, not two. There is no card for the sheet link to be on."""
    found = _gathering(monkeypatch, cards=[])

    assert not any("No sheet link on their card" in one for one in found.holes)


def test_a_board_with_no_card_and_no_receipt_says_both_plainly(monkeypatch):
    found = _gathering(monkeypatch, cards=[], receipt=("", "No invoice email for her"))

    assert any("anywhere on the board" in one for one in found.holes)
    assert any("No invoice email" in one for one in found.holes)


def test_the_timeline_carries_the_day_the_sheet_was_handed_over(monkeypatch):
    """Without it the spine is a charge and a chargeback with nothing in
    between — which is the half that answers the dispute. Juliana's timeline
    came out one line long."""
    found = _gathering(monkeypatch, cards=[HER_CARD])

    said = dict(found.timeline)
    assert "08/28/2026" in said
    assert any("delivered" in one.casefold() for one in said.values())
    assert len(found.timeline) >= 2, "still just the charge"


def test_the_handover_is_dated_from_the_comment_that_carried_the_link(monkeypatch):
    """Nicole posts the link and then says "delivered", so the link is the
    comment that dates it."""
    from wilbyte import agents
    from wilbyte.bot import jobs

    notes = [
        {"text": "delivered", "when": "2026-09-30T00:00:00Z"},
        {"text": "https://docs.google.com/spreadsheets/d/abc/edit",
         "when": "2026-08-28T18:10:00Z"},
        {"text": "takeover financial", "when": "2026-08-27T00:00:00Z"},
    ]

    assert jobs._when_delivered(notes, agents) == "2026-08-28"


def test_a_top_up_months_later_is_not_when_delivery_happened(monkeypatch):
    from wilbyte import agents
    from wilbyte.bot import jobs

    notes = [
        {"text": "https://docs.google.com/spreadsheets/d/late/edit",
         "when": "2026-11-02T00:00:00Z"},
        {"text": "https://docs.google.com/spreadsheets/d/first/edit",
         "when": "2026-08-28T18:10:00Z"},
    ]

    assert jobs._when_delivered(notes, agents) == "2026-08-28"


def test_no_sheet_comment_leaves_the_timeline_alone_rather_than_guessing(monkeypatch):
    from wilbyte import agents
    from wilbyte.bot import jobs

    assert jobs._when_delivered([{"text": "delivered", "when": "2026-08-28T00:00:00Z"}],
                                agents) == ""
    assert jobs._when_delivered([], agents) == ""


# ------------------------------------- telling a signed contract from a receipt

# PandaDoc production API is behind "Request a demo" on this account, so the
# contract arrives the way it always has: downloaded by hand and dragged into
# the command. Which makes reading it correctly the whole of the job.


def test_a_pandadoc_contract_is_a_contract_not_an_invoice():
    """A services agreement quotes a price and has payment terms in it, so
    looking for invoice words first files it as a receipt."""
    from wilbyte.bot import jobs

    said = """LEAD PURCHASE AGREEMENT
    This Agreement is entered into between Agent Lead Lab and Juliana
    Hernandez. Total amount due: $129.37, invoice to follow.
    The Client agrees not to initiate a chargeback.
    --- Completion Certificate --- Electronically signed August 28, 2026.
    Generated by PandaDoc."""

    assert jobs._what_pdf_is(said) == "contract"


def test_the_completion_certificate_alone_is_enough():
    """It carries the signing date, the reference and the audit trail, which
    is the reason the contract is worth attaching."""
    from wilbyte.bot import jobs

    assert jobs._what_pdf_is(
        "Completion Certificate\nDocument completed by all parties\nAudit trail"
    ) == "contract"


def test_a_real_invoice_is_still_an_invoice():
    from wilbyte.bot import jobs

    assert jobs._what_pdf_is(
        "PAYRA\nPayment confirmation\nInvoice #INV-17754\nTotal Paid $129.37"
    ) == "invoice"
    assert jobs._what_pdf_is("Invoice 1182\nAmount due on receipt") == "invoice"


def test_a_no_chargeback_clause_marks_the_contract():
    from wilbyte.bot import jobs

    assert jobs._what_pdf_is(
        "Terms\nThe Client waives the right to a chargeback. No-chargeback."
    ) == "contract"


def test_a_pdf_that_is_neither_is_left_alone_rather_than_guessed_at():
    """Anything RYTE cannot place goes in at the end under its own heading."""
    from wilbyte.bot import jobs

    assert jobs._what_pdf_is("a page of notes about nothing in particular") == ""
    assert jobs._what_pdf_is("") == ""


def test_the_word_amount_due_buried_in_a_contract_no_longer_wins():
    """It used to be searched over the whole document rather than the front,
    so any agreement mentioning payment terms came out an invoice."""
    from wilbyte.bot import jobs

    said = "SERVICE AGREEMENT\n" + ("filler. " * 800) + "\nAmount due: $500"

    assert jobs._what_pdf_is(said) == "contract"


# ---------------------------------- the signed contract, out of the inbox

# PandaDoc's production API is behind a sales call on this account, the
# sandbox key only reaches sandbox documents, and n8n is out — so there is no
# webhook and no public address for a Mac that sleeps. But PandaDoc emails the
# completed document with the PDF on it, and that email is the contract.


class Inbox:
    """A stubbed Gmail reader pinned to one sender."""

    def __init__(self, found, *, pdf=b"%PDF-1.4 signed"):
        self.found, self.pdf, self.asked = found, pdf, []

    def invoices_for(self, *terms, since=None):
        self.asked.append(terms)
        return list(self.found)

    def download(self, message_id, attachment_id):
        return self.pdf

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False


def _email(subject, *, files=(), body="Signed by Juliana Hernandez"):
    from wilbyte import gmail

    return gmail.Found(
        message_id="m1", subject=subject, when="Fri, 28 Aug 2026 13:00:00",
        body=body, files=list(files),
    )


def _contract(monkeypatch, emails, *, sender="pandadoc.com", pdf=b"%PDF-1.4 signed"):
    from types import SimpleNamespace

    from wilbyte import gmail, rebuttal
    from wilbyte.bot import jobs

    box = Inbox(emails, pdf=pdf)
    monkeypatch.setattr(gmail, "open_contracts", lambda secrets: box)
    config = SimpleNamespace(secrets=SimpleNamespace(gmail_contract_sender=sender))
    return jobs._signed_contract(config, rebuttal.read_facts(JULIANA)), box


def test_the_completed_document_brings_its_pdf_back(monkeypatch):
    (said, pdf, name, trouble), box = _contract(
        monkeypatch,
        [_email("Juliana Hernandez completed the document",
                files=[("Lead Purchase Agreement.pdf", "a1")])],
    )

    assert pdf == b"%PDF-1.4 signed"
    assert name == "Lead Purchase Agreement.pdf"
    assert "Juliana Hernandez" in said
    assert trouble == ""
    assert box.asked == [("Juliana Hernandez",)]


def test_the_completed_one_wins_over_sent_and_viewed(monkeypatch):
    """PandaDoc emails at every step, and only the completed one is signed."""
    (said, _, _, _), _ = _contract(monkeypatch, [
        _email("Document sent to Juliana Hernandez"),
        _email("Juliana Hernandez completed Lead Purchase Agreement",
               files=[("agreement.pdf", "a1")], body="completed"),
    ])

    assert "completed Lead Purchase Agreement" in said


def test_an_email_with_no_pdf_says_so_and_still_gives_what_it_says(monkeypatch):
    """Who signed and when is most of what the timeline wants, even with no
    document to carry the clause."""
    (said, pdf, _, trouble), _ = _contract(
        monkeypatch, [_email("Juliana Hernandez completed the document")],
    )

    assert pdf == b""
    assert "Juliana Hernandez" in said
    assert "no PDF on it" in trouble


def test_nothing_in_the_inbox_is_a_hole_rather_than_a_crash(monkeypatch):
    (said, pdf, _, trouble), _ = _contract(monkeypatch, [])

    assert (said, pdf) == ("", b"")
    assert "No signed contract" in trouble


def test_not_configured_is_not_a_hole_in_the_document(monkeypatch):
    """A hole saying "nobody set up Gmail" is about RYTE, not the dispute."""
    (said, pdf, _, trouble), _ = _contract(monkeypatch, [], sender="")

    assert (said, pdf, trouble) == ("", b"", "")


def test_the_contract_reader_is_pinned_to_its_own_sender():
    """The sender is the boundary of what can be read, set in the module so
    nothing calling it can widen it."""
    from types import SimpleNamespace

    from wilbyte import gmail

    made = gmail.open_contracts(SimpleNamespace(
        gmail_contract_sender="pandadoc.com",
        google_client_id="i", google_client_secret="s", google_refresh_token="r",
        gmail_refresh_token="", gmail_client_id="", gmail_client_secret="",
    ))

    assert made._sender == "pandadoc.com"


def test_no_contract_sender_names_its_own_setting():
    from types import SimpleNamespace

    from wilbyte import gmail

    with pytest.raises(gmail.GmailError) as raised:
        gmail.open_contracts(SimpleNamespace(
            gmail_contract_sender="",
            google_client_id="i", google_client_secret="s",
            google_refresh_token="r", gmail_refresh_token="",
        ))

    assert "GMAIL_CONTRACT_SENDER" in str(raised.value)


def test_a_found_contract_stops_the_document_asking_for_one():
    from wilbyte import rebuttal

    holes = rebuttal.what_is_missing(
        rebuttal.Gathered(contract="Juliana Hernandez completed the document"), [],
    )

    assert not any("signed contract" in one for one in holes)


def test_no_contract_anywhere_still_asks_for_it():
    from wilbyte import rebuttal

    holes = rebuttal.what_is_missing(rebuttal.Gathered(), [])

    assert any("signed contract" in one for one in holes)


# ------------------------------- the reason code decides what to argue

# Franklin's own two rebuttals, side by side. Juliana Hernandez's answers code
# 37 — the cardholder saying she never authorised the payment — with who paid:
# the AVS match on her home address, the order confirmed from the number on
# her invoice, the payment link she clicked. Jose Zambrano's answers "not as
# described", where none of that matters and the answer is the tier he chose,
# the clauses he signed, and him working the leads for months.


@pytest.mark.parametrize(
    "block, code",
    [
        ("Code: 37 - No Cardholder Authorization", "37"),
        ("Reason: Not as Described\ncode: 13.3", "13.3"),
        ("Reason Code: 13.1", "13.1"),
        ("code: 10.4", "10.4"),
    ],
)
def test_the_code_is_found_wherever_it_was_written(block, code):
    assert rebuttal.read_facts(f"Customer Name: X\n{block}").code == code


def test_a_dollar_amount_is_not_a_reason_code():
    """"$ 129.37" ends in the digits of code 37, and reading it as one aims
    the whole document at the wrong question while looking correct."""
    one = rebuttal.read_facts("Customer Name: X\nDispute Amount: $ 129.37")

    assert one.code == ""
    assert one.amount == "$129.37"


@pytest.mark.parametrize(
    "block",
    [
        "ARN: 72307626241809574244780",
        "Card Number (Last 4): 7543",
        "Transaction Date: 8/28/2026",
    ],
)
def test_the_other_numbers_on_a_notice_are_not_codes(block):
    assert rebuttal.read_facts(f"Customer Name: X\n{block}").code == ""


def test_code_37_is_aimed_at_who_paid():
    one = rebuttal.read_facts("Customer Name: X\nCode: 37")
    said = rebuttal.aimed_at(one)

    assert "37" in said and "No Cardholder Authorization" in said
    assert "WHO paid" in said
    assert "AVS" in said
    assert "concedes the point" in said


def test_not_as_described_is_aimed_at_what_arrived():
    one = rebuttal.read_facts("Customer Name: X\nCode: 13.3")
    said = rebuttal.aimed_at(one)

    assert "WHAT was promised" in said
    assert "not guaranteed" in said
    assert "AVS" not in said, "that is the other code's argument"


def test_no_code_argues_the_whole_record_rather_than_guessing():
    """The code lives in the ElevateQS portal, not on the notification."""
    said = rebuttal.aimed_at(rebuttal.read_facts("Customer Name: X"))

    assert "NO REASON CODE WAS GIVEN" in said
    assert "whole record" in said


def test_a_code_nobody_knows_is_not_invented():
    assert rebuttal.code_in("Code: 99.9") == ""
    assert rebuttal.what_the_code_means("99.9") is None


def test_the_prompt_carries_the_aim():
    one = rebuttal.read_facts("Customer Name: X\nDispute Amount: $10\nCode: 37")
    said = rebuttal.writing_prompt(one, rebuttal.Gathered(), [])

    assert "THE REASON CODE IS 37" in said


# ------------------------------- aged leads are not asked for a contract

# "aged leads we dont have contracts, only fresh/new agents. Juliana is an
# aged leads."


def test_an_aged_leads_order_is_not_asked_for_a_contract():
    holes = rebuttal.what_is_missing(rebuttal.Gathered(aged=True), [])

    assert not any("signed contract" in one for one in holes)


def test_a_new_agent_still_is():
    holes = rebuttal.what_is_missing(rebuttal.Gathered(aged=False), [])

    assert any("signed contract" in one for one in holes)


def test_the_aged_flag_comes_off_their_card(monkeypatch):
    found = _gathering(monkeypatch, cards=[HER_CARD])

    assert found.aged is True, "AGED LEAD - Juliana Hernandez"


def test_a_new_agent_card_is_not_aged(monkeypatch):
    found = _gathering(
        monkeypatch, cards=[{"id": "c1", "name": "New Agent - Steve Dass"}],
    )

    assert found.aged is False


def test_the_inbox_is_not_searched_for_a_contract_that_never_existed(monkeypatch):
    """Not just unasked for — not looked for either."""
    from wilbyte.bot import jobs

    asked = []
    monkeypatch.setattr(
        jobs, "_signed_contract",
        lambda config, dispute: (asked.append(1), ("", b"", "", ""))[1],
    )
    _gathering(monkeypatch, cards=[HER_CARD])

    assert asked == []


# ----------------------- what only the payment portal knows, pasted in

# Section 2 of Franklin's Juliana rebuttal — AVS Y, initiated by customer, the
# billing address matching, the issuer's approval — is the whole answer to
# code 37 and none of it is in the Payra receipt email or reachable by any API
# RYTE has. So it is pasted, under a line that says where it starts.


PORTAL = """Initiated By       Customer
AVS Response       Y
Billing name       Juliana Hernandez, 13722 Brownsville Street, Houston, TX 77015
Authorization      100 - approved
Transaction ID     da8ro570i476liker12g"""


@pytest.mark.parametrize(
    "marker",
    ["PAYMENT:", "payment", "GATEWAY:", "Transaction details:", "Invoice log:"],
)
def test_the_record_starts_where_it_says_it_does(marker):
    block, paid = rebuttal.split_payment(
        f"Customer Name: Juliana Hernandez\n{marker}\n{PORTAL}"
    )

    assert paid == PORTAL
    assert "Customer Name" in block


def test_the_portal_lines_are_not_parsed_as_dispute_fields():
    """"AVS Response  Y" read as a labelled field is how an email address
    becomes the letter Y."""
    block, _ = rebuttal.split_payment(
        f"Customer Name: Juliana Hernandez\nCustomer Email: hjuliana650@gmail.com\n"
        f"PAYMENT:\n{PORTAL}"
    )
    one = rebuttal.read_facts(block)

    assert one.customer_email == "hjuliana650@gmail.com"
    assert one.customer_name == "Juliana Hernandez"
    assert "Brownsville" not in one.customer_name


def test_no_marker_leaves_everything_as_the_dispute_block():
    """Every rebuttal before this one worked without one."""
    block, paid = rebuttal.split_payment("Customer Name: X\nDispute Amount: $10")

    assert paid == ""
    assert rebuttal.read_facts(block).customer_name == "X"


def test_the_prompt_carries_the_gateway_record_and_says_what_it_means():
    one = rebuttal.read_facts("Customer Name: X\nDispute Amount: $10\nCode: 37")
    said = rebuttal.writing_prompt(one, rebuttal.Gathered(payment=PORTAL), [])

    assert "THE GATEWAY'S OWN RECORD" in said
    assert "AVS Response       Y" in said
    assert "customer-initiated means" in said


def test_nothing_pasted_adds_no_section():
    one = rebuttal.read_facts("Customer Name: X\nDispute Amount: $10")
    said = rebuttal.writing_prompt(one, rebuttal.Gathered(), [])

    assert "THE GATEWAY'S OWN RECORD" not in said


# ------------------------- a screenshot of the portal does the same job

# "yeah ill just send the screenshots for that for you to attach" — RYTE
# already transcribes what is in an attached image, so the gateway record
# reaches the document the same way whether it was pasted or photographed.


def test_a_payment_record_is_one_of_the_kinds_of_exhibit():
    assert "payment" in rebuttal.EXHIBITS
    assert "AVS" in rebuttal.EXHIBITS["payment"]


def test_it_is_filed_next_to_the_invoice_not_at_the_back():
    """"other" goes in at the end under its own heading, which is where a
    gateway record would have landed."""
    assert rebuttal.UNDER["payment"] == rebuttal.UNDER["invoice"] + 1


def test_the_exhibits_are_lettered_with_it_in_that_order():
    made = [
        rebuttal.Exhibit(name="texts.png", kind="texts"),
        rebuttal.Exhibit(name="gateway.png", kind="payment"),
        rebuttal.Exhibit(name="invoice.pdf", kind="invoice"),
    ]
    lettered = rebuttal.letter_them(made)

    by_kind = {one.kind: one.letter for one in lettered}
    assert by_kind["invoice"] < by_kind["payment"]


def test_a_transcribed_screenshot_reaches_the_writing():
    """The transcript is what carries the AVS value, not the caption."""
    shot = rebuttal.Exhibit(
        name="gateway.png", kind="payment",
        caption="payment gateway record",
        transcript="Initiated By  Customer\nAVS Response  Y",
    )
    lettered = rebuttal.letter_them([shot])
    one = rebuttal.read_facts("Customer Name: X\nDispute Amount: $10\nCode: 37")
    said = rebuttal.writing_prompt(one, rebuttal.Gathered(), lettered)

    assert "AVS Response  Y" in said
    assert "Payment Authorization Record" in said


def test_pasting_it_still_wins_over_a_screenshot(monkeypatch):
    """What was typed is what somebody chose; a transcript is RYTE reading."""
    import inspect

    from wilbyte.bot import client as bot_client

    source = inspect.getsource(bot_client._build_rebuttal) if hasattr(
        bot_client, "_build_rebuttal"
    ) else ""
    # The wiring: the pasted block is only replaced when there wasn't one.
    assert "if not paid_with:" in inspect.getsource(bot_client)


# ------------------- replying to RYTE's own flag card is how you ask

# "@Ryte code: 37 — No Cardholder Authorization rebuttal", in reply to the
# chargeback card RYTE had just posted, came back "I need a bit more of the
# dispute notice". Every fact was in the message being replied to — RYTE's
# own — and it could not read its own handwriting.

FLAG_CARD = """⚖ Chargeback
• Customer — Juliana Hernandez
• Amount — $129.37
• Transaction — 8/28/2026
• ARN — 72307626241809574244780"""


def test_ryte_can_read_its_own_flag_card():
    one = rebuttal.read_facts(FLAG_CARD)

    assert one.missing() == []
    assert one.customer_name == "Juliana Hernandez"
    assert one.amount == "$129.37"
    assert one.transaction_date == "8/28/2026"
    assert one.arn == "72307626241809574244780"


def test_the_portal_notice_still_reads_the_same():
    """The bullet form is additional, not a replacement."""
    one = rebuttal.read_facts(JULIANA)

    assert one.customer_name == "Juliana Hernandez"
    assert one.customer_email == "hjuliana650@gmail.com"
    assert one.transaction_date == "8/28/2026"
    assert one.missing() == []
    assert rebuttal.read_facts(
        JULIANA + "\nCode: 37 - No Cardholder Authorization"
    ).code == "37"


def test_customer_email_is_not_claimed_by_the_name():
    """"Customer" now matches the name field, and it must let "Customer
    Email" past to the one below it."""
    one = rebuttal.read_facts(
        "Customer Email: hjuliana650@gmail.com\nCustomer Name: Juliana Hernandez"
    )

    assert one.customer_email == "hjuliana650@gmail.com"
    assert one.customer_name == "Juliana Hernandez"


def test_a_transaction_id_is_not_a_transaction_date():
    one = rebuttal.read_facts("Transaction ID: da8ro570i476liker12g")

    assert one.transaction_date == ""


def test_a_dash_in_prose_is_not_a_labelled_field():
    """The bullet is what makes it a field. Prose is full of dashes."""
    one = rebuttal.read_facts("the leads — all 25 of them — arrived on time")

    assert one.missing() == ["Customer Name", "Dispute Dollar Amount",
                             "Transaction Date"]


# ------------------------------------- finding the notice without being replied to


HER_FLAG_CARD = """⚖️ **Chargeback**
• **Customer** — Juliana Hernandez
• **Amount** — $129.37
• **Transaction** — 8/28/2026
• **ARN** — 7230762624180957424478
[The notice](https://discord.com/x)"""

HER_NOTICE = """Disputed Payment | Elevateqs ❌

ARN: 7230762624180957424478
Customer Name: Juliana Hernandez
Customer Email: hjuliana650@gmail.com
Card Number (Last 4): 7543
Transaction Date: 8/28/2026
Dispute Amount: $ 129.37"""


def _asked(said, above, monkeypatch, *, reference=None):
    """Run the rebuttal handler with `above` sitting in the channel already."""
    import asyncio

    from wilbyte.bot import client as bot_client
    from wilbyte.bot import jobs

    spoke = []

    class Responder:
        requester_id = 1

        async def send(self, content=None, **kwargs):
            spoke.append(str(content or kwargs.get("embed") or ""))

    class Older:
        attachments: list = []
        jump_url = "https://discord.com/above"

        def __init__(self, content):
            self.content = content

    class Channel:
        async def history(self, limit=0, before=None):  # pragma: no cover - shape
            raise NotImplementedError

    class Message:
        attachments: list = []
        channel = Channel()

    # discord.py's history() is an async iterator, not a coroutine.
    async def history(limit=0, before=None):
        for one in above:
            yield Older(one)

    Message.channel.history = history
    Message.reference = reference

    def stop(*args, **kwargs):
        raise TypeError("far enough")

    monkeypatch.setattr(jobs, "rebuttal_evidence", stop)
    monkeypatch.setattr(bot_client.embeds, "error", lambda text, **kw: f"ERROR: {text}")

    class Config:
        class secrets:
            anthropic_api_key = "x"

    asyncio.run(bot_client._rebuttal(Responder(), Config(), Message(), said))
    return spoke


def test_the_code_on_its_own_reads_the_notice_already_in_the_channel(monkeypatch):
    """Franklin typed "@RYTE code: 37 - No Cardholder Authorization rebuttal"
    in the dispute channel, with the notice and RYTE's own flag card a few
    lines above it, and RYTE asked him to paste a block that was on the screen
    twice. Both parse whole - it just never looked."""
    spoke = _asked(
        "code: 37 — No Cardholder Authorization rebuttal",
        [HER_FLAG_CARD, HER_NOTICE],
        monkeypatch,
    )

    assert not any("I need a bit more" in one for one in spoke), spoke
    assert any("Juliana Hernandez" in one for one in spoke), spoke


def test_the_code_he_typed_survives_the_notice_it_reads_back(monkeypatch):
    """Neither the notice nor the flag card carries the reason code - it is on
    the ElevateQS portal. Taking the fuller message must not drop the half he
    added, or the rebuttal argues the wrong question."""
    from wilbyte.bot import client as bot_client
    from wilbyte import rebuttal as rules

    class Older:
        content = HER_FLAG_CARD

    fuller = bot_client._fuller_dispute(
        rules, rules.read_facts("code: 37"), "code: 37", "", Older()
    )

    assert fuller is not None
    assert fuller[0].code == "37"
    assert fuller[0].customer_name == "Juliana Hernandez"


def test_a_message_nearby_cannot_overwrite_what_he_typed(monkeypatch):
    """A fuller message wins; a thinner one is left alone. Otherwise the last
    dispute in the channel quietly replaces the one being asked about."""
    from wilbyte.bot import client as bot_client
    from wilbyte import rebuttal as rules

    class Older:
        content = "Customer Name: Someone Else"

    whole = rules.read_facts(HER_NOTICE)

    assert bot_client._fuller_dispute(rules, whole, HER_NOTICE, "", Older()) is None


def test_it_says_which_message_the_facts_came_off(monkeypatch):
    """Reading back is a guess about what somebody meant. A rebuttal built for
    the wrong customer is not a thing to find inside the finished document."""
    spoke = _asked(
        "code: 37 — No Cardholder Authorization rebuttal",
        [HER_FLAG_CARD],
        monkeypatch,
    )

    assert any("discord.com/above" in one for one in spoke), spoke


def test_a_reason_code_line_is_not_a_customer_name():
    """`named_in` read "code" out of "code: 37 - No Cardholder Authorization"
    and addressed the rebuttal to a customer called code."""
    from wilbyte import rebuttal as rules

    assert rules.named_in("code: 37 — No Cardholder Authorization rebuttal") == ""
    assert rules.named_in("rebuttal Jose Zambrano") == "Jose Zambrano"


# ----------------------------------- the gaps Juliana's first rebuttal didn't name


def test_a_who_paid_code_asks_for_the_gateway_record():
    """Juliana's rebuttal was built under code 37 and named one gap - texts.
    It argued who paid from the invoice alone, so the acquirer was asked to
    take our word for the AVS match and the billing name captured at payment,
    neither of which was in front of them."""
    from wilbyte import rebuttal as rules

    holes = rules.what_is_missing(
        rules.Gathered(), [], rules.read_facts("code: 37\nCustomer Name: Juliana")
    )

    assert any("authorisation record" in one for one in holes), holes


def test_a_pasted_payment_record_closes_that_gap():
    from wilbyte import rebuttal as rules

    found = rules.Gathered()
    found.payment = "AVS: Y  Customer initiated: yes  Approval: 004411"
    holes = rules.what_is_missing(found, [], rules.read_facts("code: 37"))

    assert not any("authorisation record" in one for one in holes), holes


def test_a_not_as_described_code_does_not_ask_for_it():
    """On 13.3 the authorisation data is beside the point, and a gap named
    that does not matter is one that gets ignored along with the ones that do."""
    from wilbyte import rebuttal as rules

    holes = rules.what_is_missing(rules.Gathered(), [], rules.read_facts("code: 13.3"))

    assert not any("authorisation record" in one for one in holes), holes


def test_no_dispute_date_is_named_because_the_timeline_loses_a_line():
    """Juliana's notice carried no dispute date, so the timeline ended at the
    charge and nobody was told why."""
    from wilbyte import rebuttal as rules

    holes = rules.what_is_missing(
        rules.Gathered(), [], rules.read_facts("Transaction Date: 8/28/2026")
    )

    assert any("date the chargeback was filed" in one for one in holes), holes


def test_a_dispute_date_that_is_there_is_not_asked_for():
    from wilbyte import rebuttal as rules

    holes = rules.what_is_missing(
        rules.Gathered(), [], rules.read_facts("Dispute Date: 9/8/2026")
    )

    assert not any("date the chargeback was filed" in one for one in holes), holes


def test_the_gaps_are_still_named_with_no_dispute_at_all():
    """`what_is_missing` is called from two places and the dispute is new to
    it. Neither caller may break for want of it."""
    from wilbyte import rebuttal as rules

    assert rules.what_is_missing(rules.Gathered(), []) is not None


# ------------------------------- the tracker tab has two tables on it


SEPT_ROW_ONE = [
    "Name of Disputer", "Date of Transaction", "Amount", "Closer", "Status",
    "", "Sep 2026 chargebacks — deductions by closer",
]


def test_only_the_lists_own_columns_are_read():
    """The month tab holds the disputes in A-E and a deductions-by-closer
    summary from column G. Reading all of row 1 made a seven-wide row out of
    five headings, an empty column F and the summary's title - and the preview
    offered a column called "(unnamed)"."""
    from wilbyte.bot import jobs

    assert jobs._the_list_columns(SEPT_ROW_ONE) == [
        "Name of Disputer", "Date of Transaction", "Amount", "Closer", "Status",
    ]


def test_the_row_goes_under_the_list_not_under_the_other_table():
    """Google's append was handed the tab and chose the summary, writing
    Juliana Hernandez below it in the summary's columns - under headings that
    every one of them meant something else."""
    from wilbyte.bot import jobs

    asked = []

    class Reading:
        def rows(self, sheet, span):
            asked.append(span)
            return [SEPT_ROW_ONE[:5], ["a"] * 5, ["b"] * 5, ["c"] * 5]

    where, at = jobs._a_free_row(Reading(), "sid", "Sept 2026", 5)

    assert asked == ["'Sept 2026'!A:E"]
    assert (where, at) == ("'Sept 2026'!A5:E5", 5)


def test_the_free_row_is_read_rather_than_assumed():
    """`put` overwrites. A dispute written over another dispute is the one
    thing that must never happen to somebody's tracker."""
    from wilbyte.bot import jobs

    class Reading:
        def rows(self, sheet, span):
            return [["h"]] + [["x"]] * 40

    where, at = jobs._a_free_row(Reading(), "sid", "Sept 2026", 5)

    assert at == 42 and where == "'Sept 2026'!A42:E42"


@pytest.mark.parametrize(
    "wide, letter", [(1, "A"), (5, "E"), (7, "G"), (26, "Z"), (27, "AA")]
)
def test_the_last_column_is_named_correctly(wide, letter):
    from wilbyte.bot import jobs

    assert jobs._column_letter(wide) == letter


def test_the_tracker_write_says_which_row_it_landed_on():
    """A write nobody can check is a write nobody trusts - and this one went
    to the wrong table once already."""
    from wilbyte.bot import jobs

    written = {}

    class Writing:
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

        def rows(self, sheet, span):
            return [["h"] * 5, ["a"] * 5]

        def put(self, sheet, span, rows):
            written["span"] = span
            return span

        def tabs(self, sheet):
            return [{"title": "Sept 2026", "sheetId": 7}]

        def match_row_above(self, sheet, tab_id, row, wide):
            written["styled"] = (tab_id, row, wide)

        def append(self, *a, **kw):  # pragma: no cover - must not be used
            raise AssertionError("append picks the wrong table on this sheet")

    from wilbyte import gsheets

    class Config:
        class secrets:
            tracker_sheet_id = "sid"

    import unittest.mock as mock

    with mock.patch.object(gsheets, "SheetsClient", lambda creds: Writing()), \
            mock.patch.object(gsheets, "credentials", lambda s: None):
        where, trouble = jobs.track_chargeback(Config(), "Sept 2026", ["a"] * 5)

    assert not trouble, trouble
    assert written["span"] == "'Sept 2026'!A3:E3"
    assert where == "Sept 2026, row 3"
    assert written["styled"] == (7, 3, 5)


# ------------------------------- what Juliana's first draft got wrong in prose


def _prompt():
    from wilbyte import rebuttal as rules

    return rules.writing_prompt(rules.read_facts("code: 37"), rules.Gathered(), [])


def test_two_counts_of_the_same_thing_have_to_be_reconciled():
    """Her rebuttal said "25 leads requested" in the summary and "35 delivered"
    in section 4, in the same document, unexplained."""
    said = _prompt()

    assert "two different counts" in said and "more were delivered" in said


def test_words_are_only_put_in_the_cardholders_mouth_when_she_wrote_them():
    """It listed lead names - "Alma R. Aguirre", "Hector anguiano" - among the
    notes she typed. If those are delivered rows rather than her handwriting,
    the cardholder contradicts the sentence and the paragraph goes with it."""
    said = _prompt()

    assert "not their handwriting" in said
    assert "unmistakably theirs" in said


def test_our_own_records_are_not_described_as_thin():
    """"our record of that delivery is the word delivered alongside the sheet
    link" is an opinion about our evidence, volunteered to the acquirer."""
    said = _prompt()

    assert "never how much of it there is" in said


# ------------------------------------- the document's shape, against his own


def test_the_subtitle_names_the_code_being_answered():
    """"Response to Cardholder Dispute" says nothing the title has not. The
    acquirer's reader sorts these by code."""
    from wilbyte import rebuttal as rules

    assert rules.answering(rules.read_facts("code: 37")) == (
        "AGENT LEAD LAB — Response to Reason Code 37, No Cardholder Authorization"
    )


def test_an_unknown_code_does_not_get_invented_into_the_subtitle():
    from wilbyte import rebuttal as rules

    assert rules.answering(rules.read_facts("Customer Name: X")).endswith(
        "Response to Cardholder Dispute"
    )


def test_the_fact_table_says_what_was_bought():
    """A reader deciding a dispute wants the product in the first ten seconds,
    not in paragraph four of the summary."""
    from wilbyte import rebuttal as rules

    rows = dict(rules.header(
        rules.read_facts("Customer Name: X\nDispute Amount: $1"),
        rules.Gathered(product="25 Aged Final Expense — Texas"),
    ))

    assert rows["Product"] == "25 Aged Final Expense — Texas"


def test_the_fact_table_still_builds_with_nothing_gathered():
    """`header` is called from two places and `found` is new to it."""
    from wilbyte import rebuttal as rules

    assert rules.header(rules.read_facts("Customer Name: X"))


def test_a_set_of_fields_is_written_as_a_table_not_a_sentence():
    """Franklin's own rebuttal for this dispute is four tables, and they are
    why it can be checked field by field in the time an acquirer has."""
    said = _prompt()

    assert "TABLE: Field | Value | What it means" in said
    assert "first line without one ends the table" in said


def test_a_table_is_read_back_out_of_the_writing():
    from wilbyte import rebuttaldoc

    assert rebuttaldoc._TABLE.match("TABLE: Field | Value | What it means")
    assert rebuttaldoc._cells("Initiated By | **Customer** | started it") == [
        "Initiated By", "Customer", "started it",
    ]


def test_a_numbered_column_is_not_given_two_inches_of_white_space():
    """"# | Fact" has one digit in its first column and a sentence squeezed
    into its second."""
    from wilbyte import rebuttaldoc

    numbered = rebuttaldoc._column_widths(["#", "Fact"], [["1", "..."]], 2)
    labelled = rebuttaldoc._column_widths(
        ["Field", "Value"], [["Initiated By", "Customer"]], 2
    )

    assert numbered[0] == 0.55
    assert sum(numbered) == pytest.approx(sum(labelled))


def test_the_tables_reach_the_finished_file(tmp_path):
    """Written and rendered are two different things, and the prompt has been
    right while the document was wrong before."""
    from wilbyte import rebuttal as rules, rebuttaldoc
    import docx

    said = {"body": """SUMMARY:
She paid it herself.

ARGUMENT: Payment authorization data (Exhibit B)
The gateway record establishes who paid.
TABLE: Field | Value | What it means
Initiated By | Customer | The cardholder started the payment.
AVS Response | Y | Full match on the billing address.

CONCLUSION:
The record answers the code.
TABLE: # | Fact
1 | Customer-initiated, AVS Y.
2 | Paid through a link she clicked.
""", "messages": ""}

    where = tmp_path / "r.docx"
    rebuttaldoc.build(
        rules.read_facts("Customer Name: Juliana Hernandez\nDispute Amount: $129.37\n"
                         "Transaction Date: 8/28/2026\ncode: 37"),
        said, rules.Gathered(), [], into=where,
    )
    made = docx.Document(str(where))
    shapes = [[c.text.strip() for c in row.cells] for t in made.tables for row in t.rows]

    assert ["Field", "Value", "What it means"] in shapes
    assert ["Initiated By", "Customer", "The cardholder started the payment."] in shapes
    assert ["#", "Fact"] in shapes
    # And the prose around them is still prose.
    assert any("The gateway record establishes who paid." == p.text.strip()
               for p in made.paragraphs)


def test_a_line_of_prose_with_a_pipe_in_it_is_not_swallowed(tmp_path):
    """A table ends at the first line with no pipe. Only lines under a TABLE:
    are read as rows, so a stray pipe in a sentence cannot start one."""
    from wilbyte import rebuttal as rules, rebuttaldoc
    import docx

    where = tmp_path / "r.docx"
    rebuttaldoc.build(
        rules.read_facts("Customer Name: X\nDispute Amount: $1\nTransaction Date: 8/28/2026"),
        {"body": "SUMMARY:\nThe invoice reads Paid | Settled, as shown.", "messages": ""},
        rules.Gathered(), [], into=where,
    )
    made = docx.Document(str(where))

    assert any("Paid | Settled" in p.text for p in made.paragraphs)


def test_blanks_are_filled_from_the_channel_but_nothing_is_overwritten():
    """The flag card answers `missing()`, so reading stopped there and the
    notice below it carrying her email, her card's last four and the MID was
    never opened. Those rows belong on the fact table.

    Filled, never replaced: an older dispute further up the channel must not
    be able to rewrite the facts of the one being answered."""
    from wilbyte.bot import client as bot_client
    from wilbyte import rebuttal as rules

    class Older:
        content = HER_NOTICE + "\nMID: 510200014664"

    one = rules.read_facts(HER_FLAG_CARD)
    one.customer_email = "typed-by-hand@example.com"

    bot_client._fill_the_blanks(rules, one, Older())

    assert one.mid == "510200014664"
    assert one.card == "7543"
    assert one.customer_email == "typed-by-hand@example.com"
    assert one.customer_name == "Juliana Hernandez"


def test_the_handler_fills_the_fact_table_from_the_whole_channel(monkeypatch):
    """End to end: the code typed, the flag card above it, the notice above
    that. The finished fact table needs all three."""
    import asyncio

    from wilbyte.bot import client as bot_client
    from wilbyte.bot import jobs

    seen = {}

    class Responder:
        requester_id = 1

        async def send(self, content=None, **kwargs):
            return None

    class Older:
        attachments: list = []
        jump_url = "https://discord.com/above"

        def __init__(self, content):
            self.content = content

    class Message:
        attachments: list = []
        reference = None

        class channel:
            @staticmethod
            async def history(limit=0, before=None):
                for one in (HER_FLAG_CARD, HER_NOTICE + "\nMID: 510200014664"):
                    yield Older(one)

    def stop(config, dispute):
        seen["dispute"] = dispute
        raise TypeError("far enough")

    monkeypatch.setattr(jobs, "rebuttal_evidence", stop)
    monkeypatch.setattr(bot_client.embeds, "error", lambda text, **kw: text)

    class Config:
        class secrets:
            anthropic_api_key = "x"

    asyncio.run(bot_client._rebuttal(
        Responder(), Config(), Message(), "code: 37 — No Cardholder Authorization"
    ))

    one = seen["dispute"]

    assert one.customer_name == "Juliana Hernandez"
    assert one.code == "37"
    assert one.mid == "510200014664"
    assert one.card == "7543"
    assert one.customer_email == "hjuliana650@gmail.com"
